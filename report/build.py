"""Build the project report: report/GrammarGuard_Project_Report.docx (editable) and an HTML file that is printed to PDF.

    python report/build.py            # writes the DOCX and report/_report.html
"""

from __future__ import annotations

import ast
import base64
import html
import re
import sys
import textwrap
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
HERE = ROOT / "report"
sys.path.insert(0, str(HERE))
import content as C  # noqa: E402

BASE_ROLE = {
    "willard": "Regex to finite-state machine to an index of allowed vocabulary tokens: the design of our token index.",
    "geng": "Grammar-constrained decoding of structured outputs: basis of our grammar (LL(1)) decoding.",
    "koo": "Automata-based constraints (DFA and pushdown) with token-level lookup: closest to our approach.",
    "thompson": "Regular expression to epsilon-NFA construction (nfa.py).",
    "rabin": "Subset construction, NFA to DFA (dfa.py).",
    "kleene": "Regular expressions equal finite automata; DFA to regex by state elimination (elimination.py).",
    "myhill": "Myhill-Nerode theorem: the basis of DFA minimization.",
    "nerode": "Myhill-Nerode theorem: the basis of DFA minimization.",
    "lewis": "Top-down (LL) parsing: basis of the LL(1) table and the pushdown automaton (cfg.py, pda.py).",
    "hmu": "Primary textbook for the automata and grammar algorithms.",
}
REF_NO = {r[0]: i + 1 for i, r in enumerate(C.REFS)}

# ------------------------------------------------------------------ figure / listing numbering
FIG: dict[str, int] = {}
for b in C.B:
    if b[0] == "img":
        FIG[b[4]] = len(FIG) + 1


# ------------------------------------------------------------------ inline markup
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[A-Za-z][^*\n]*?\*|(?:\{\{[a-z0-9]+\}\})+|\{fig:[a-z0-9_]+\})")


def inline(text: str):
    """Split text into (kind, value) pieces: text, bold, italic, code, cite (list of numbers), fig (id)."""
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            for k2, v2 in inline(part[2:-2]):  # code (and citations) inside bold stay code; plain text becomes bold
                out.append(("bold" if k2 == "text" else k2, v2))
        elif part.startswith("`"):
            out.append(("code", part[1:-1]))
        elif part.startswith("{{"):
            keys = re.findall(r"\{\{([a-z0-9]+)\}\}", part)
            out.append(("cite", sorted(REF_NO[k] for k in keys)))
        elif part.startswith("{fig:"):
            out.append(("fig", part[5:-1]))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(("italic", part[1:-1]))
        else:
            out.append(("text", part))
    return out


def plain(s: str) -> str:
    """Heading text without inline markup characters."""
    return s.replace("`", "")


def cite_text(nums):
    return "[" + ", ".join(str(n) for n in nums) + "]"


# ------------------------------------------------------------------ code extraction (verbatim from the repository)
def function_source(relpath: str, qualname: str) -> str:
    src = (ROOT / relpath).read_text()
    tree = ast.parse(src)
    parts = qualname.split(".")
    scope = tree.body
    node = None
    for i, name in enumerate(parts):
        node = next(n for n in scope if getattr(n, "name", None) == name)
        scope = getattr(node, "body", [])
    seg = ast.get_source_segment(src, node)
    first_line = src.splitlines()[node.lineno - 1]
    indent = len(first_line) - len(first_line.lstrip())
    return textwrap.dedent(" " * indent + seg).rstrip()


LISTING: dict[str, int] = {}
n = 0
for b in C.B:
    if b[0] in ("code", "codefn"):
        n += 1
        LISTING[b[2] if b[0] == "codefn" else b[2]] = n


# ================================================================== DOCX
def build_docx(out: Path):
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    TEXT_W = 6.5  # inches usable

    NAVY = RGBColor(0x1F, 0x38, 0x64)

    def set_font(style, name, size, bold=None, color=None, italic=None):
        style.font.name = name
        style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if italic is not None:
            style.font.italic = italic
        if color is not None:
            style.font.color.rgb = color
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), name)

    set_font(doc.styles["Normal"], "Times New Roman", 11.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.12
    set_font(doc.styles["Heading 1"], "Calibri", 17, True, NAVY)
    set_font(doc.styles["Heading 2"], "Calibri", 13.5, True, NAVY)
    set_font(doc.styles["Heading 3"], "Calibri", 12, True, RGBColor(0x33, 0x41, 0x55), italic=False)
    for h, before, after in (("Heading 1", 0, 8), ("Heading 2", 12, 4), ("Heading 3", 8, 3)):
        pf = doc.styles[h].paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(before), Pt(after), True

    # ---- helpers
    def shade(cell_or_par_pr, fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        cell_or_par_pr.append(shd)

    def runs(par, text, size=None, in_table=False, color=None, bold_all=False):
        for kind, val in inline(text):
            if kind == "cite":
                r = par.add_run(cite_text(val))
            elif kind == "fig":
                r = par.add_run(str(FIG[val]) if in_table else f"Figure {FIG[val]}")
            else:
                r = par.add_run(val)
                if kind == "bold":
                    r.bold = True
                elif kind == "italic":
                    r.italic = True
                elif kind == "code":
                    r.font.name = "Consolas"
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                    r.font.size = Pt((size or 11.5) - 1.5)
                    r.font.color.rgb = RGBColor(0x9A, 0x1B, 0x4B)
            if size and kind != "code":
                r.font.size = Pt(size)
            if bold_all:
                r.bold = True
            if color is not None:
                r.font.color.rgb = color
        return par

    def para(text, align=None, size=None, space_after=None, style=None):
        p = doc.add_paragraph(style=style)
        runs(p, text, size=size)
        if align:
            p.alignment = align
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def field(par, instr, placeholder=""):
        for t, extra in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
            r = par.add_run()
            if t == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = extra
                r._r.append(el)
            elif t == "text":
                r.text = extra
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), t)
                r._r.append(el)

    def table(header, rows, widths):
        total = sum(widths)
        t = doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblpr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblpr.append(lay)
        w_in = [TEXT_W * w / total for w in widths]
        for ci, h in enumerate(header):
            c = t.rows[0].cells[ci]
            c.width = Inches(w_in[ci])
            shade(c._tc.get_or_add_tcPr(), "1F3864")
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            runs(p, h, size=9.5, in_table=True, color=RGBColor(255, 255, 255), bold_all=True)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                c = t.rows[ri].cells[ci]
                c.width = Inches(w_in[ci])
                if ri % 2 == 0:
                    shade(c._tc.get_or_add_tcPr(), "F1F5F9")
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                runs(p, val, size=9.5, in_table=True)
            cant = OxmlElement("w:cantSplit")
            cant.set(qn("w:val"), "true")
            t.rows[ri]._tr.get_or_add_trPr().append(cant)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        before = len(p.runs)
        runs(p, text, size=9.5)
        for r in p.runs[before:]:
            if r.font.name != "Consolas":
                r.italic = True
        p.paragraph_format.space_after = Pt(10)
        return p

    def code_block(src):
        p = doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        shade(ppr, "F3F4F6")
        pf = p.paragraph_format
        pf.space_after, pf.line_spacing, pf.left_indent = Pt(8), 1.0, Cm(0.15)
        lines = src.replace("\t", "    ").split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(7.8)
            if i < len(lines) - 1:
                r.add_break()

    NO_BREAK = {"2. Project Title", "4. Register Numbers", "5. Students' Names"}

    # ---- footer with page number
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    tabs = fp.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(TEXT_W), alignment=2)  # right
    r = fp.add_run(f"{C.SHORT_TITLE}: Project Report\t")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    r2 = fp.add_run()
    r2.font.size = Pt(9)
    field(fp, "PAGE", "1")
    sec.different_first_page_header_footer = True  # no footer on the cover

    # ---- body
    for b in C.B:
        kind = b[0]
        if kind == "cover":
            for _ in range(2):
                doc.add_paragraph()
            p = para(f"**{C.COURSE.upper()}**", WD_ALIGN_PARAGRAPH.CENTER, 13, 4)
            para("Project Report", WD_ALIGN_PARAGRAPH.CENTER, 13, 28)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(C.SHORT_TITLE)
            r.bold = True
            r.font.size = Pt(34)
            r.font.name = "Calibri"
            r.font.color.rgb = NAVY
            p.paragraph_format.space_after = Pt(6)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(C.TITLE.split(": ", 1)[1])
            r.font.size = Pt(15)
            r.font.name = "Calibri"
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            p.paragraph_format.space_after = Pt(40)
            para("**Submitted by**", WD_ALIGN_PARAGRAPH.CENTER, 12, 6)
            for name, reg in C.TEAM:
                para(f"{name}  ({reg})", WD_ALIGN_PARAGRAPH.CENTER, 13, 2)
            doc.add_paragraph()
            para(f"{C.DATE}", WD_ALIGN_PARAGRAPH.CENTER, 12, 14)
            para(f"Source code: {C.REPO}", WD_ALIGN_PARAGRAPH.CENTER, 10.5, 0)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "toc":
            p = doc.add_paragraph()
            r = p.add_run("Table of Contents")
            r.bold = True
            r.font.size = Pt(17)
            r.font.name = "Calibri"
            r.font.color.rgb = NAVY
            tp = doc.add_paragraph()
            field(tp, 'TOC \\o "1-2" \\h \\z \\u', "Right-click here and choose Update Field to build the table of contents.")
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "h1":
            h = doc.add_heading(b[1], level=1)
            if b[1] not in NO_BREAK and b[1] != "1. Aim of the Project":
                h.paragraph_format.page_break_before = True
        elif kind == "h2":
            doc.add_heading(plain(b[1]), level=2)
        elif kind == "h3":
            doc.add_heading(plain(b[1]), level=3)
        elif kind == "p":
            p = para(b[1])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "title_box":
            p = doc.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            shade(ppr, "EEF2FF")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b[1])
            r.bold = True
            r.font.size = Pt(15)
            r.font.name = "Calibri"
            r.font.color.rgb = NAVY
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
        elif kind == "ul":
            for item in b[1]:
                p = doc.add_paragraph(style="List Bullet")
                runs(p, item)
                p.paragraph_format.space_after = Pt(3)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "ol":  # numbered by hand so that every list starts at 1 (Word would continue one numbering)
            for i, item in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(3)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.add_run(f"{i}.  ")
                runs(p, item)
        elif kind == "table":
            table(b[1], b[2], b[3])
        elif kind == "img":
            _, fname, cap, width, fid = b
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(str(HERE / "img" / fname), width=Inches(min(width, TEXT_W)))
            caption(f"Figure {FIG[fid]}: {cap}")
        elif kind == "code":
            caption_text = f"Listing {LISTING[b[2]]}: {b[2]}"
            cp = doc.add_paragraph()
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_after = Pt(2)
            r = cp.add_run(caption_text)
            r.italic = True
            r.font.size = Pt(9.5)
            code_block(b[1])
        elif kind == "codefn":
            _, relpath, qual, expl = b
            cp = doc.add_paragraph()
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_after = Pt(2)
            r = cp.add_run(f"Listing {LISTING[qual]}: ")
            r.bold = True
            r.font.size = Pt(10)
            runs(cp, f"`{relpath}`, `{qual}`. {expl}", size=10)
            code_block(function_source(relpath, qual))
        elif kind == "refs":
            build_refs_docx(doc, runs, table, shade, caption)
        else:
            raise ValueError(kind)

    # ask Word to refresh the TOC / fields when the file is opened
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)
    doc.core_properties.title = C.TITLE
    doc.core_properties.author = " and ".join(n for n, _ in C.TEAM)
    doc.core_properties.subject = f"{C.COURSE} project report"
    doc.save(out)
    return out


def build_refs_docx(doc, runs, table, shade, caption):
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement  # noqa: F401

    p = doc.add_paragraph()
    runs(p, "References marked **★** are the **base papers** of this project; they are shaded and summarised in Section 17.1. The **base links** (open-source implementations, model card and documentation the work builds on) are listed in Section 17.3. Numbers in square brackets in the text refer to this list. Every title, author list, year and DOI below was checked against the publisher or arXiv record.")
    doc.add_heading("17.1 Base papers and their role in this project", level=2)
    rows = []
    for key, grp, base, text, link in C.REFS:
        if base:
            rows.append([f"**★ [{REF_NO[key]}]**", BASE_ROLE.get(key, "")])
    table(["Base paper", "How it is used"], rows, [1.2, 8])
    doc.add_heading("17.2 Complete reference list", level=2)
    for grp, title in C.GROUP_TITLES.items():
        h = doc.add_paragraph()
        r = h.add_run(title)
        r.bold = True
        r.font.size = Pt(11)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.keep_with_next = True
        for key, g, base, text, link in C.REFS:
            if g != grp:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(26)
            p.paragraph_format.first_line_indent = Pt(-26)
            p.paragraph_format.space_after = Pt(4)
            if base:
                ppr = p._p.get_or_add_pPr()
                shade(ppr, "FEF3C7")
            r = p.add_run(f"[{REF_NO[key]}]  " + ("★ " if base else ""))
            r.bold = True
            r.font.size = Pt(10)
            r = p.add_run(text + (" " if link else ""))
            r.font.size = Pt(10)
            if link:
                r = p.add_run(link)
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    doc.add_heading("17.3 Base links and software documentation", level=2)
    rows = []
    for label, url, note, base in C.LINKS:
        rows.append([("**★ " + label + "**") if base else label, url, note])
    table(["Resource", "Link", "Role"], rows, [3, 4.2, 3.4])


# ================================================================== HTML (for PDF)
CSS = """
@page { size: A4; margin: 2.1cm 2.1cm 2.3cm 2.1cm;
  @bottom-left { content: "GrammarGuard: Project Report"; font: 9pt Calibri, Arial, sans-serif; color: #64748b; }
  @bottom-right { content: counter(page); font: 9pt Calibri, Arial, sans-serif; color: #64748b; } }
@page :first { @bottom-left { content: ""; } @bottom-right { content: ""; } }
* { box-sizing: border-box; }
body { font: 11pt/1.42 "Times New Roman", Times, serif; color: #111; margin: 0; }
h1 { font: 700 17pt Calibri, Arial, sans-serif; color: #1F3864; border-bottom: 2px solid #1F3864; padding-bottom: 3pt; margin: 0 0 10pt; break-before: page; break-after: avoid; }
h1.nb, h1.first { break-before: auto; margin-top: 18pt; }
h2 { font: 700 13.5pt Calibri, Arial, sans-serif; color: #1F3864; margin: 14pt 0 4pt; break-after: avoid; }
h3 { font: 700 12pt Calibri, Arial, sans-serif; color: #334155; margin: 10pt 0 3pt; break-after: avoid; }
p { margin: 0 0 6pt; text-align: justify; }
ul, ol { margin: 0 0 6pt 0; padding-left: 20pt; } li { margin-bottom: 3pt; text-align: justify; }
code { font: 9pt Consolas, "Cascadia Mono", monospace; color: #9a1b4b; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0 10pt; font: 9.5pt/1.25 Calibri, Arial, sans-serif; }
th code { color: inherit; } th { background: #1F3864; color: #fff; text-align: left; padding: 4pt 5pt; border: 1px solid #1F3864; }
td { border: 1px solid #94a3b8; padding: 3pt 5pt; vertical-align: top; } tr { break-inside: avoid; } tbody tr:nth-child(even) td { background: #F1F5F9; }
td code { font-size: 8.6pt; }
figure { margin: 8pt 0 10pt; text-align: center; break-inside: avoid; }
figure img { border: 1px solid #cbd5e1; max-width: 100%; }
figcaption { font: italic 9.5pt Calibri, Arial, sans-serif; margin-top: 3pt; }
pre { background: #F3F4F6; border-left: 3px solid #1F3864; padding: 6pt 8pt; font: 7.6pt/1.25 Consolas, "Cascadia Mono", monospace; white-space: pre-wrap; overflow-wrap: anywhere; margin: 0 0 9pt; }
.listing { font: 10pt Calibri, Arial, sans-serif; margin: 8pt 0 2pt; break-after: avoid; }
.titlebox { background: #EEF2FF; border: 1px solid #c7d2fe; text-align: center; font: 700 15pt Calibri, Arial, sans-serif; color: #1F3864; padding: 10pt; margin: 4pt 0 10pt; }
.cover { text-align: center; padding-top: 70pt; break-after: page; }
.cover .course { font: 700 13pt Calibri, Arial, sans-serif; letter-spacing: 1px; } .cover .big { font: 700 36pt Calibri, Arial, sans-serif; color: #1F3864; margin: 24pt 0 4pt; }
.cover .sub { font: 15pt Calibri, Arial, sans-serif; color: #334155; margin-bottom: 50pt; } .cover p { text-align: center; margin: 3pt; }
.toc { break-after: page; } .toc h1 { break-before: auto; } .toc div { font: 11pt Calibri, Arial, sans-serif; margin: 1pt 0; } .toc .l2 { margin-left: 18pt; font-size: 10pt; color: #475569; }
.refs p { text-align: left; padding-left: 26pt; text-indent: -26pt; font: 9.8pt/1.3 Calibri, Arial, sans-serif; margin-bottom: 4pt; } .refs p.base { background: #FEF3C7; }
.refs a, .link { color: #1D4ED8; word-break: break-all; text-decoration: none; }
.refgroup { font: 700 11pt Calibri, Arial, sans-serif; margin: 10pt 0 3pt; break-after: avoid; }
"""


def esc(s):
    return html.escape(s, quote=False)


def inline_html(text, in_table=False):
    out = []
    for kind, val in inline(text):
        if kind == "bold":
            out.append(f"<b>{esc(val)}</b>")
        elif kind == "italic":
            out.append(f"<i>{esc(val)}</i>")
        elif kind == "code":
            out.append(f"<code>{esc(val)}</code>")
        elif kind == "cite":
            out.append(esc(cite_text(val)))
        elif kind == "fig":
            out.append(str(FIG[val]) if in_table else f"Figure {FIG[val]}")
        else:
            out.append(esc(val))
    return "".join(out)


def data_uri(path: Path) -> str:
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode()


def build_html(out: Path):
    NO_BREAK = {"2. Project Title", "4. Register Numbers", "5. Students' Names"}
    parts = []
    heads = [(b[0], b[1]) for b in C.B if b[0] in ("h1", "h2")]
    for b in C.B:
        k = b[0]
        if k == "cover":
            parts.append(f"""<div class="cover"><div class="course">{C.COURSE.upper()}</div><p style="font:13pt Calibri">Project Report</p>
<div class="big">{C.SHORT_TITLE}</div><div class="sub">{esc(C.TITLE.split(': ', 1)[1])}</div>
<p><b>Submitted by</b></p>{''.join(f'<p style="font-size:13pt">{esc(n)} &nbsp;({r})</p>' for n, r in C.TEAM)}
<p style="margin-top:30pt">{C.DATE}</p><p style="font-size:10.5pt">Source code: {C.REPO}</p></div>""")
        elif k == "toc":
            items = "".join(f'<div class="{"l2" if t == "h2" else ""}">{esc(txt)}</div>' for t, txt in heads)
            parts.append(f'<div class="toc"><h1 class="nb">Table of Contents</h1>{items}</div>')
        elif k == "h1":
            cls = "nb" if b[1] in NO_BREAK or b[1] == "1. Aim of the Project" else ""
            parts.append(f'<h1 class="{cls}">{esc(b[1])}</h1>')
        elif k in ("h2", "h3"):
            parts.append(f"<{k}>{esc(plain(b[1]))}</{k}>")
        elif k == "p":
            parts.append(f"<p>{inline_html(b[1])}</p>")
        elif k == "title_box":
            parts.append(f'<div class="titlebox">{esc(b[1])}</div>')
        elif k in ("ul", "ol"):
            parts.append(f"<{k}>" + "".join(f"<li>{inline_html(i)}</li>" for i in b[1]) + f"</{k}>")
        elif k == "table":
            _, header, rows, widths = b
            tot = sum(widths)
            cols = "".join(f'<col style="width:{100 * w / tot:.1f}%">' for w in widths)
            head = "".join(f"<th>{inline_html(h, True)}</th>" for h in header)
            body = "".join("<tr>" + "".join(f"<td>{inline_html(c, True)}</td>" for c in row) + "</tr>" for row in rows)
            parts.append(f"<table><colgroup>{cols}</colgroup><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
        elif k == "img":
            _, fname, cap, width, fid = b
            parts.append(f'<figure><img src="{data_uri(HERE / "img" / fname)}" style="width:{min(width, 6.5) / 6.5 * 100:.0f}%"><figcaption>Figure {FIG[fid]}: {inline_html(cap)}</figcaption></figure>')
        elif k == "code":
            parts.append(f'<div class="listing"><i>Listing {LISTING[b[2]]}: {esc(b[2])}</i></div><pre>{esc(b[1])}</pre>')
        elif k == "codefn":
            _, relpath, qual, expl = b
            parts.append(f'<div class="listing"><b>Listing {LISTING[qual]}:</b> <code>{esc(relpath)}</code>, <code>{esc(qual)}</code>. {esc(expl)}</div><pre>{esc(function_source(relpath, qual))}</pre>')
        elif k == "refs":
            parts.append(refs_html())
    doc = f'<!doctype html><html><head><meta charset="utf-8"><title>{esc(C.TITLE)}</title><style>{CSS}</style></head><body>{"".join(parts)}</body></html>'
    out.write_text(doc, encoding="utf-8")
    return out


def refs_html():
    o = ['<div class="refs">']
    o.append("<p style='padding-left:0;text-indent:0;text-align:justify;font:10.5pt Calibri'>References marked <b>★</b> are the <b>base papers</b> of this project; they are shaded and summarised in Section 17.1. The <b>base links</b> (open-source implementations, model card and documentation the work builds on) are listed in Section 17.3. Numbers in square brackets in the text refer to this list. Every title, author list, year and DOI below was checked against the publisher or arXiv record.</p>")
    o.append("<h2>17.1 Base papers and their role in this project</h2><table><colgroup><col style='width:12%'><col style='width:88%'></colgroup><thead><tr><th>Base paper</th><th>How it is used</th></tr></thead><tbody>")
    for key, grp, base, text, link in C.REFS:
        if base:
            o.append(f"<tr><td><b>★ [{REF_NO[key]}]</b></td><td>{esc(BASE_ROLE.get(key, ''))}</td></tr>")
    o.append("</tbody></table><h2>17.2 Complete reference list</h2>")
    for grp, title in C.GROUP_TITLES.items():
        o.append(f'<div class="refgroup">{esc(title)}</div>')
        for key, g, base, text, link in C.REFS:
            if g == grp:
                a = f' <a class="link" href="{html.escape(link)}">{esc(link)}</a>' if link else ""
                o.append(f'<p class="{"base" if base else ""}"><b>[{REF_NO[key]}] {"★ " if base else ""}</b>{esc(text)}{a}</p>')
    o.append("<h2>17.3 Base links and software documentation</h2><table><colgroup><col style='width:30%'><col style='width:42%'><col style='width:28%'></colgroup><thead><tr><th>Resource</th><th>Link</th><th>Role</th></tr></thead><tbody>")
    for label, url, note, base in C.LINKS:
        name = f"<b>★ {esc(label)}</b>" if base else esc(label)
        o.append(f'<tr><td>{name}</td><td><a class="link" href="{html.escape(url)}">{esc(url)}</a></td><td>{esc(note)}</td></tr>')
    o.append("</tbody></table></div>")
    return "".join(o)


if __name__ == "__main__":
    d = build_docx(HERE / "GrammarGuard_Project_Report.docx")
    h = build_html(HERE / "_report.html")
    print("wrote", d, f"({d.stat().st_size // 1024} KB)")
    print("wrote", h, f"({h.stat().st_size // 1024} KB)")
